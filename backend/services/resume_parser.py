"""
backend/services/resume_parser.py
Multi-format resume parsing engine supporting PDF, DOCX, legacy binary DOC, and TXT/MD.
Provides format decoding, Unicode NFKC normalization, heuristic section segmentation,
contact entity regex extraction, and optional Gemini AI structured extraction.
"""

import io
import re
import uuid
import logging
import unicodedata
from typing import Tuple, Dict, Any, List, Optional

import pypdf
import docx
from docx.text.paragraph import Paragraph
from docx.table import Table

from backend.models.resume_schemas import (
    ParsedResume,
    CandidateContact,
    EducationEntry,
    WorkExperienceEntry,
    ProjectEntry,
    CandidateSkills,
)

logger = logging.getLogger(__name__)


# -------------------------------------------------------------
# Custom Exception Hierarchy
# -------------------------------------------------------------

class ResumeParsingError(Exception):
    """Base exception for all resume parsing failures."""
    pass


class UnsupportedFormatError(ResumeParsingError):
    """Raised when an unsupported file extension or MIME type is uploaded."""
    pass


class CorruptedFileError(ResumeParsingError):
    """Raised when a document is damaged, unreadable, or invalid format."""
    pass


class PasswordProtectedFileError(ResumeParsingError):
    """Raised when an encrypted or password-protected PDF is encountered."""
    pass


class ScannedOrEmptyResumeError(ResumeParsingError):
    """Raised when an uploaded resume contains no extractable text or is a scanned image."""
    pass


# -------------------------------------------------------------
# Pre-Compiled Regex & Taxonomy Constants
# -------------------------------------------------------------

EMAIL_REGEX = re.compile(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,7}\b")
PHONE_REGEX = re.compile(r"(?:\+?\d{1,3}[-.\s]?)?\(?\d{3,4}\)?[-.\s]?\d{3,4}[-.\s]?\d{3,4}\b")
GITHUB_REGEX = re.compile(r"(?:https?:\/\/)?(?:www\.)?github\.com\/([A-Za-z0-9_.-]+(?:\/[A-Za-z0-9_.-]+)?)", re.I)
LINKEDIN_REGEX = re.compile(r"(?:https?:\/\/)?(?:www\.)?linkedin\.com\/in\/([A-Za-z0-9_-]+)", re.I)
PORTFOLIO_REGEX = re.compile(
    r"https?:\/\/(?:[a-zA-Z0-9-]+\.)*(?:vercel\.app|netlify\.app|onrender\.com|github\.io|[a-zA-Z0-9-]+\.(?:io|dev|me))(?:\/[^\s,;()>]*)?",
    re.I
)

SECTION_PATTERNS = {
    "summary": re.compile(r"^(?:#+\s*)?(?:professional\s+summary|career\s+objective|summary|profile|about\s+me|objective)[:\s]*$", re.I),
    "education": re.compile(r"^(?:#+\s*)?(?:education|academic\s+background|academics|educational\s+qualifications|qualifications)[:\s]*$", re.I),
    "skills": re.compile(r"^(?:#+\s*)?(?:technical\s+skills|core\s+competencies|skills\s*(?:&|and)\s*tools|technologies|skills|technical\s+proficiencies)[:\s]*$", re.I),
    "experience": re.compile(r"^(?:#+\s*)?(?:work\s+experience|professional\s+experience|employment\s+history|experience|internships)[:\s]*$", re.I),
    "projects": re.compile(r"^(?:#+\s*)?(?:personal\s+projects|academic\s+projects|featured\s+projects|key\s+projects|projects|open\s+source)[:\s]*$", re.I),
    "certifications": re.compile(r"^(?:#+\s*)?(?:certifications|certificates|licenses|achievements|honors\s*(?:&|and)\s*awards)[:\s]*$", re.I),
}

TECH_KEYWORDS = {
    "python", "javascript", "typescript", "go", "golang", "rust", "java", "c++", "c#", "ruby", "php",
    "fastapi", "flask", "django", "react", "next.js", "vue", "angular", "node.js", "express", "spring",
    "postgresql", "postgres", "mysql", "mongodb", "redis", "sqlite", "elasticsearch", "cassandra",
    "docker", "kubernetes", "aws", "gcp", "azure", "git", "linux", "ci/cd", "graphql", "rest api",
    "pytest", "junit", "html", "css", "tailwind", "sass", "redux", "kafka", "rabbitmq", "celery",
    "terraform", "ansible", "pytorch", "tensorflow", "scikit-learn", "pandas", "numpy", "grpc", "tokio", "webassembly"
}

TOOL_KEYWORDS = {
    "docker", "kubernetes", "git", "github", "gitlab", "jira", "postman", "vscode", "pycharm",
    "linux", "bash", "powershell", "terraform", "jenkins", "github actions", "grafana", "prometheus",
    "figma", "notion", "slack", "webpack", "vite", "npm", "yarn", "pip", "uv"
}

SOFT_KEYWORDS = {
    "agile", "scrum", "leadership", "communication", "problem solving", "teamwork",
    "mentoring", "cross-functional", "project management", "collaboration", "critical thinking",
    "ownership", "time management", "adaptability", "client interaction"
}


# -------------------------------------------------------------
# Parser Service Class
# -------------------------------------------------------------

class ResumeParserService:
    """Multi-format resume parsing and structured profile extraction service."""

    @staticmethod
    def normalize_text(raw_text: str) -> str:
        """
        Normalizes extracted document text:
        - Resolves Unicode ligatures and canonical forms (NFKC)
        - Normalizes Windows CRLF and classic Mac CR to standard LF
        - Strips zero-width characters and BOM
        - Replaces non-breaking spaces with standard spaces
        - Canonicalizes smart quotes, em-dashes, and en-dashes
        - Collapses repeated blank lines
        """
        if not raw_text:
            return ""

        # NFKC normalizes ligatures: fi, fl, ffi, etc.
        text = unicodedata.normalize("NFKC", raw_text)
        # Canonicalize line endings
        text = text.replace("\r\n", "\n").replace("\r", "\n")
        # Remove zero-width spaces and BOM
        text = re.sub(r"[\u200b\ufeff]", "", text)
        # Normalize non-breaking space
        text = text.replace("\u00a0", " ")
        # Normalize quotes and dashes
        text = text.replace("“", '"').replace("”", '"').replace("‘", "'").replace("’", "'")
        text = text.replace("—", "-").replace("–", "-")
        # Collapse multiple horizontal whitespace (except newlines)
        text = re.sub(r"[ \t]+", " ", text)
        # Collapse excessive newlines
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()

    @classmethod
    def extract_text_from_pdf(cls, content: bytes) -> str:
        """
        Extracts text from PDF bytes via pypdf.PdfReader.
        Handles encrypted PDFs and detects scanned/image documents.
        """
        if not content:
            raise ScannedOrEmptyResumeError("The uploaded resume contains no extractable text or is corrupted/scanned.")

        try:
            stream = io.BytesIO(content)
            reader = pypdf.PdfReader(stream)
        except Exception as e:
            raise CorruptedFileError(f"The uploaded PDF is invalid or corrupted: {str(e)}")

        # Check encryption
        if reader.is_encrypted:
            try:
                decrypt_result = reader.decrypt("")
                if decrypt_result == 0:
                    raise PasswordProtectedFileError("Password-protected PDF files cannot be parsed.")
            except PasswordProtectedFileError:
                raise
            except Exception:
                raise PasswordProtectedFileError("Password-protected PDF files cannot be parsed.")

        page_texts: List[str] = []
        empty_pages = 0
        total_pages = len(reader.pages)

        for page in reader.pages:
            try:
                txt = page.extract_text() or ""
                if "/Annots" in page:
                    for annot in page["/Annots"]:
                        try:
                            obj = annot.get_object()
                            if "/A" in obj and "/URI" in obj["/A"]:
                                uri = str(obj["/A"]["/URI"]).strip()
                                if uri:
                                    txt += f"\n{uri}"
                        except Exception:
                            pass
                if not txt.strip():
                    empty_pages += 1
                page_texts.append(txt)
            except Exception as e:
                logger.warning("Error extracting page in PDF: %s", e)
                empty_pages += 1

        full_text = "\n".join(page_texts).strip()
        is_scanned = len(full_text) < 40 or (total_pages > 0 and empty_pages == total_pages)

        if len(full_text) == 0 or (is_scanned and len(full_text.strip()) < 40):
            raise ScannedOrEmptyResumeError("The uploaded resume contains no extractable text or is corrupted/scanned.")

        return full_text

    @classmethod
    def extract_text_from_docx(cls, content: bytes) -> str:
        """
        Extracts text from DOCX bytes via python-docx.
        Iterates element.body to preserve exact visual document order between
        paragraphs and embedded tables. Deduplicates merged table cells.
        """
        if not content:
            raise ScannedOrEmptyResumeError("The uploaded resume contains no extractable text or is corrupted/scanned.")

        try:
            stream = io.BytesIO(content)
            doc = docx.Document(stream)
        except Exception as e:
            raise CorruptedFileError(f"Failed to parse DOCX document: invalid or corrupted archive: {str(e)}")

        items: List[str] = []

        # Extract headers from sections (often contains contact details)
        for section in doc.sections:
            try:
                for hp in section.header.paragraphs:
                    htxt = hp.text.strip()
                    if htxt and htxt not in items:
                        items.append(htxt)
            except Exception:
                pass

        # Iterate body in document flow order
        try:
            for block in doc.element.body:
                tag = block.tag.split("}")[-1]
                if tag == "p":
                    p = Paragraph(block, doc)
                    ptxt = p.text.strip()
                    if ptxt:
                        items.append(ptxt)
                elif tag == "tbl":
                    tbl = Table(block, doc)
                    for row in tbl.rows:
                        row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        # Deduplicate sequential duplicate cells resulting from table cell merges
                        deduped: List[str] = []
                        for t in row_texts:
                            if not deduped or deduped[-1] != t:
                                deduped.append(t)
                        if deduped:
                            items.append(" | ".join(deduped))
        except Exception:
            # Fallback to standard paragraphs and tables if body iteration fails
            for p in doc.paragraphs:
                if p.text.strip():
                    items.append(p.text.strip())
            for tbl in doc.tables:
                for row in tbl.rows:
                    r_text = " | ".join(dict.fromkeys(c.text.strip() for c in row.cells if c.text.strip()))
                    if r_text:
                        items.append(r_text)

        full_text = "\n".join(items).strip()
        if len(full_text) < 20:
            raise ScannedOrEmptyResumeError("The uploaded resume contains no extractable text or is corrupted/scanned.")

        return full_text

    @classmethod
    def extract_text_from_doc(cls, content: bytes) -> str:
        """
        Extracts text from legacy binary Word (.doc) files via OLE2 stream analysis.
        Extracts UTF-16LE and ASCII text streams and filters internal Word keywords.
        """
        if not content:
            raise ScannedOrEmptyResumeError("The uploaded resume contains no extractable text or is corrupted/scanned.")

        is_ole2 = content.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")
        if not is_ole2:
            raise CorruptedFileError("The uploaded legacy Word document lacks a valid OLE2 CFBF signature.")

        # Extract UTF-16LE text sequences (2 bytes per char: printable ASCII followed by \x00)
        utf16_pattern = re.compile(b"(?:[\x20-\x7E\t\n\r]\x00){4,}")
        utf16_matches = [m.group(0).decode("utf-16-le", errors="ignore") for m in utf16_pattern.finditer(content)]

        # Extract ASCII/CP1252 printable sequences
        ascii_pattern = re.compile(b"[\x20-\x7E\t\n\r]{4,}")
        ascii_matches = [m.group(0).decode("latin-1", errors="ignore") for m in ascii_pattern.finditer(content)]

        utf16_text = "\n".join(s.strip() for s in utf16_matches if len(s.strip()) > 3)
        ascii_text = "\n".join(s.strip() for s in ascii_matches if len(s.strip()) > 3)

        combined_lines: List[str] = []
        for block in (ascii_text, utf16_text):
            for line in block.split("\n"):
                cleaned = line.strip()
                if cleaned and cleaned not in combined_lines:
                    combined_lines.append(cleaned)

        boilerplate = {
            "worddocument", "summaryinformation", "documentsummaryinformation",
            "compobj", "normal.dotm", "normal.dot", "microsoft word", "times new roman",
            "calibri", "arial", "courier new", "symbol", "wingdings", "heading 1", "heading 2",
            "table grid", "default paragraph font"
        }

        lines: List[str] = []
        for line in combined_lines:
            cleaned = line.strip()
            if cleaned and cleaned.lower() not in boilerplate and not all(c in "=-_#* " for c in cleaned):
                lines.append(cleaned)

        full_text = "\n".join(lines).strip()
        if len(full_text) < 20:
            raise ScannedOrEmptyResumeError("The uploaded resume contains no extractable text or is corrupted/scanned.")

        return full_text

    @classmethod
    def extract_text_from_txt(cls, content: bytes) -> str:
        """Extracts text from plain text or markdown resumes (.txt, .md)."""
        if not content:
            raise ScannedOrEmptyResumeError("The uploaded resume contains no extractable text or is corrupted/scanned.")

        try:
            full_text = content.decode("utf-8")
        except UnicodeDecodeError:
            try:
                full_text = content.decode("latin-1")
            except Exception as e:
                raise CorruptedFileError(f"Invalid text encoding: {str(e)}")

        if len(full_text.strip()) < 20:
            raise ScannedOrEmptyResumeError("The uploaded resume contains no extractable text or is corrupted/scanned.")

        return full_text.strip()

    @classmethod
    def parse_document(cls, filename: str, content: bytes) -> str:
        """
        Public contract matching PROJECT.md:
        ResumeParserService.parse_document(filename: str, content: bytes) -> str
        Extracts and normalizes text from PDF, DOCX, DOC, TXT, or MD documents.
        """
        fn_lower = filename.lower() if filename else ""

        if fn_lower.endswith(".pdf") or content.startswith(b"%PDF"):
            raw_text = cls.extract_text_from_pdf(content)
        elif fn_lower.endswith(".docx") or content.startswith(b"PK\x03\x04"):
            raw_text = cls.extract_text_from_docx(content)
        elif fn_lower.endswith(".doc") or content.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
            raw_text = cls.extract_text_from_doc(content)
        elif fn_lower.endswith(".txt") or fn_lower.endswith(".md"):
            raw_text = cls.extract_text_from_txt(content)
        else:
            ext = f".{filename.split('.')[-1]}" if "." in filename else "unknown"
            raise UnsupportedFormatError(f"Unsupported file format '{ext}'. Supported formats: .pdf, .docx, .doc, .txt, .md.")

        normalized = cls.normalize_text(raw_text)
        if len(normalized) < 20:
            raise ScannedOrEmptyResumeError("The uploaded resume contains no extractable text or is corrupted/scanned.")

        return normalized

    @classmethod
    def segment_sections(cls, text: str) -> Dict[str, str]:
        """
        Segments normalized resume text into structured sections:
        header, summary, education, skills, experience, projects, certifications, other.
        """
        lines = text.split("\n")
        current_sec = "header"
        sections: Dict[str, List[str]] = {
            "header": [],
            "summary": [],
            "education": [],
            "skills": [],
            "experience": [],
            "projects": [],
            "certifications": [],
            "other": [],
        }

        for line in lines:
            cleaned = line.strip().strip("#*-_:").strip()
            matched_sec = None
            if len(cleaned) < 45:
                for sec_name, pattern in SECTION_PATTERNS.items():
                    if pattern.match(cleaned):
                        matched_sec = sec_name
                        break
            if matched_sec:
                current_sec = matched_sec
            else:
                sections[current_sec].append(line)

        return {k: "\n".join(v).strip() for k, v in sections.items() if v}

    @classmethod
    def extract_structured_candidate(
        cls,
        raw_text: str,
        filename: Optional[str] = None,
        candidate_id: Optional[str] = None,
        gemini_api_key: Optional[str] = None,
    ) -> ParsedResume:
        """
        Extracts structured ParsedResume from normalized raw text.
        Dual Mode:
        - If gemini_api_key is provided and valid, invokes Gemini AI (gemini-2.5-flash).
        - If gemini_api_key is missing or API call fails, falls back to deterministic regex heuristics.
        """
        cid = candidate_id or f"cand_{uuid.uuid4().hex[:10]}"

        # Mode A: Gemini AI Assisted Extraction
        if gemini_api_key and gemini_api_key.strip():
            try:
                from google import genai
                from google.genai import types

                client = genai.Client(api_key=gemini_api_key.strip())
                prompt = (
                    "You are the ProofCard Resume Entity Extractor. Parse the following resume text "
                    "into the structured ParsedResume JSON schema.\n\n"
                    f"Candidate ID: {cid}\n"
                    f"Filename: {filename or 'resume'}\n\n"
                    "Resume Text:\n"
                    f"{raw_text[:8000]}"
                )
                response = client.models.generate_content(
                    model="gemini-2.5-flash",
                    contents=prompt,
                    config=types.GenerateContentConfig(
                        response_mime_type="application/json",
                        response_schema=ParsedResume,
                    ),
                )
                if response.text:
                    parsed = ParsedResume.model_validate_json(response.text)
                    parsed.candidate_id = cid
                    parsed.raw_text = raw_text
                    if filename:
                        parsed.file_metadata["filename"] = filename
                    return parsed
            except Exception as e:
                logger.warning("Gemini AI extraction failed (%s), falling back to heuristic parsing.", e)

        # Mode B: Deterministic Heuristic Extractor
        return cls._extract_heuristic(raw_text, filename=filename, candidate_id=cid)

    @classmethod
    def _extract_heuristic(
        cls,
        raw_text: str,
        filename: Optional[str] = None,
        candidate_id: Optional[str] = None,
    ) -> ParsedResume:
        """Rule-based, offline heuristic extraction of candidate entities."""
        cid = candidate_id or f"cand_{uuid.uuid4().hex[:10]}"
        sections = cls.segment_sections(raw_text)
        lines = [l.strip() for l in raw_text.split("\n") if l.strip()]

        # 1. Candidate Name
        generic_headers = {"resume", "curriculum vitae", "cv", "profile", "bio", "contact"}
        name = "Anonymous Candidate"
        for line in lines[:5]:
            c = line.strip()
            if c and c.lower() not in generic_headers and len(c) < 50 and not re.search(r"[@\(\)\d|:]", c):
                name = c.strip("#*-_ ")
                break

        # 2. Contact Details
        email_match = EMAIL_REGEX.search(raw_text)
        email = email_match.group(0).lower() if email_match else None

        phone = None
        for p in PHONE_REGEX.findall(raw_text):
            if not re.match(r"^\d{4}\s*-\s*\d{4}$", p.strip()):
                phone = p.strip()
                break

        all_github_urls = []
        for m in GITHUB_REGEX.finditer(raw_text):
            val = f"https://github.com/{m.group(1)}" if m.group(1) else m.group(0)
            if val not in all_github_urls:
                all_github_urls.append(val)
        github_url = all_github_urls[0] if all_github_urls else None

        li_match = LINKEDIN_REGEX.search(raw_text)
        linkedin_url = f"https://linkedin.com/in/{li_match.group(1)}" if li_match else None

        all_live_urls = []
        for m in PORTFOLIO_REGEX.finditer(raw_text):
            val = m.group(0)
            if val not in all_live_urls and "github.com" not in val.lower():
                all_live_urls.append(val)

        candidate_domains = re.findall(
            r'\b([a-zA-Z0-9-]+\.(?:com|org|dev|io|me|app|tech|ai|site|net|co|in|agency|store))\b',
            raw_text,
            re.IGNORECASE
        )
        skip_domains = {
            'gmail.com', 'yahoo.com', 'hotmail.com', 'outlook.com', 'example.com',
            'example.edu', 'example.org', 'github.com', 'linkedin.com', 'b.tech', 'm.tech', 'b.sc', 'm.sc'
        }
        for d in candidate_domains:
            d_lower = d.lower()
            if d_lower not in skip_domains and not re.search(r'[\w.+-]+@' + re.escape(d), raw_text, re.IGNORECASE):
                u = f"https://{d}"
                if not any(d_lower in existing.lower() for existing in all_live_urls):
                    all_live_urls.append(u)

        portfolio_url = all_live_urls[0] if all_live_urls else None

        all_urls = re.findall(r"https?:\/\/[^\s,;()\"'>]+", raw_text)
        other_links = [u for u in all_urls if u not in (github_url, linkedin_url, portfolio_url)][:5]

        contact = CandidateContact(
            name=name,
            email=email,
            phone=phone,
            linkedin_url=linkedin_url,
            github_url=github_url,
            portfolio_url=portfolio_url,
            other_links=other_links,
            all_github_urls=all_github_urls,
            all_live_urls=all_live_urls,
        )

        # 3. Summary
        summary = sections.get("summary") or None

        # 4. Education
        education_entries: List[EducationEntry] = []
        edu_text = sections.get("education", "")
        if edu_text:
            edu_lines = [l.strip() for l in edu_text.split("\n") if l.strip()]
            inst = "Affiliated Engineering Institute"
            deg = None
            year = None
            gpa = None

            for el in edu_lines:
                years = [int(y) for y in re.findall(r"\b(19\d\d|20\d\d)\b", el)]
                if years and not year:
                    year = max(years)

                gpa_m = re.search(
                    r"\b(?:cgpa|gpa|grade|score)[\s:]*([0-9](?:\.[0-9]{1,2})?(?:\s*\/\s*(?:10|4)(?:\.0)?)?|\b[0-9]{2}(?:\.[0-9]+)?%)",
                    el,
                    re.I,
                )
                if gpa_m and not gpa:
                    gpa = gpa_m.group(1).strip()

                deg_m = re.search(r"\b(B\.?Tech|B\.?E\.?|B\.?S\.?|M\.?Tech|M\.?S\.?|Bachelor|Master|Diploma|Ph\.?D)\b.*", el, re.I)
                if deg_m and not deg:
                    deg = deg_m.group(0).split("|")[0].strip()

                if any(k in el.lower() for k in ["university", "institute", "college", "school", "academy", "technology"]):
                    clean_inst = el.split("|")[0].strip()
                    if not deg or clean_inst != deg:
                        inst = clean_inst

            education_entries.append(
                EducationEntry(
                    institution=inst,
                    degree=deg,
                    graduation_year=year,
                    grade_or_gpa=gpa,
                )
            )

        # 5. Skills
        skill_text = sections.get("skills", "") or raw_text
        tokens = set(re.findall(r"\b[a-zA-Z\+\#\.\-]+(?:\s+[a-zA-Z]+)?\b", skill_text.lower()))

        matched_tech = [k for k in sorted(TECH_KEYWORDS) if k in tokens or re.search(rf"\b{re.escape(k)}\b", skill_text, re.I)]
        matched_tools = [k for k in sorted(TOOL_KEYWORDS) if k in tokens or re.search(rf"\b{re.escape(k)}\b", skill_text, re.I)]
        matched_soft = [k for k in sorted(SOFT_KEYWORDS) if k in tokens or re.search(rf"\b{re.escape(k)}\b", skill_text, re.I)]

        skills = CandidateSkills(
            technical=matched_tech,
            tools=matched_tools,
            non_technical=matched_soft,
        )

        # 6. Projects
        project_entries: List[ProjectEntry] = []
        proj_text = sections.get("projects", "")
        if proj_text:
            blocks = [b.strip() for b in re.split(r"\n\s*\n", proj_text) if b.strip()]
            for b in blocks[:5]:
                p_lines = [l.strip() for l in b.split("\n") if l.strip()]
                if not p_lines:
                    continue
                first_line = p_lines[0]
                title = re.sub(r"\(?https?:\/\/github\.com\/[^\s)]+\)?|\(?github\.com\/[^\s)]+\)?", "", first_line).strip(" -:|()")

                r_m = GITHUB_REGEX.search(b)
                r_url = f"https://github.com/{r_m.group(1)}" if r_m else None

                l_m = PORTFOLIO_REGEX.search(b)
                l_url = l_m.group(0) if l_m else None

                desc = " ".join(l.lstrip("-*• ").strip() for l in p_lines[1:]) if len(p_lines) > 1 else first_line
                p_tech = [t for t in TECH_KEYWORDS if re.search(rf"\b{re.escape(t)}\b", b, re.I)][:6]

                project_entries.append(
                    ProjectEntry(
                        title=title or "Key Project",
                        description=desc,
                        technologies=p_tech,
                        repo_url=r_url,
                        live_url=l_url,
                    )
                )

        # 7. Work Experience
        exp_entries: List[WorkExperienceEntry] = []
        exp_text = sections.get("experience", "")
        if exp_text:
            exp_blocks = [b.strip() for b in re.split(r"\n\s*\n", exp_text) if b.strip()]
            for b in exp_blocks[:5]:
                e_lines = [l.strip() for l in b.split("\n") if l.strip()]
                if not e_lines:
                    continue
                header_line = e_lines[0]
                company = header_line.split("|")[0].split("-")[0].strip() or "Acme Corp"
                title = header_line.split("|")[-1].strip() if "|" in header_line else "Software Engineer"

                dates = re.findall(r"\b(present|\d{4}|(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\s*\d{4})\b", b, re.I)
                start_date = dates[0] if dates else None
                end_date = dates[1] if len(dates) > 1 else ("Present" if "present" in b.lower() else None)

                highlights = [l.lstrip("-*• ").strip() for l in e_lines[1:] if l.strip()]

                exp_entries.append(
                    WorkExperienceEntry(
                        company=company,
                        title=title,
                        start_date=start_date,
                        end_date=end_date,
                        is_current="present" in (end_date or "").lower(),
                        highlights=highlights,
                    )
                )

        # 8. Certifications
        cert_entries: List[str] = []
        cert_text = sections.get("certifications", "")
        if cert_text:
            cert_entries = [l.lstrip("-*• ").strip() for l in cert_text.split("\n") if l.strip()][:10]

        file_metadata = {
            "filename": filename or "resume",
            "word_count": len(raw_text.split()),
            "char_count": len(raw_text),
            "sections_detected": [k for k, v in sections.items() if v],
        }

        return ParsedResume(
            candidate_id=cid,
            contact_info=contact,
            summary=summary,
            education=education_entries,
            work_experience=exp_entries,
            projects=project_entries,
            skills=skills,
            certifications=cert_entries,
            raw_text=raw_text,
            file_metadata=file_metadata,
        )
